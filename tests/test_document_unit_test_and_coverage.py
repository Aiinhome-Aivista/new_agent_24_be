import os
import docx
import pytest
from app.tools.document_generator.docx_generator import generate_docx_evidence
from app.tools.document_generator.generator import render_autonomous_evidence_html


def test_docx_and_html_contain_unit_test_code_and_ac_coverage(tmp_path):
    out_dir = str(tmp_path)
    evidence_data = {
        "evidence_key": "EVID-DOC-COV-01",
        "traceability_id": "TRC-TEST-COV-99",
        "story": {
            "external_key": "SBP-101",
            "title": "Support Ticket Management System",
            "acceptance_criteria": [
                {"ac_key": "AC-01", "text": "POST /api/tickets with valid parameters creates ticket and returns 201"},
                {"ac_key": "AC-02", "text": "Missing mandatory fields return 400 Bad Request with field errors"},
            ]
        },
        "target_host": "http://localhost:5001",
        "collection_name": "Tickets API",
        "summary_recommendation": "API conforms",
        "decision_status": "Ready for Approval",
        "decision_summary": "All tests passed with 100% Acceptance Criteria coverage.",
        "coverage_matrix": [
            {
                "ac_key": "AC-01",
                "requirement": "POST /api/tickets with valid parameters creates ticket and returns 201",
                "covered": True,
                "test_case_keys": ["TC-SBP101-001"]
            },
            {
                "ac_key": "AC-02",
                "requirement": "Missing mandatory fields return 400 Bad Request with field errors",
                "covered": True,
                "test_case_keys": ["TC-SBP101-002"]
            }
        ],
        "coverage_report": {
            "total_acceptance_criteria": 2,
            "covered_acceptance_criteria": 2,
            "coverage_pct": 100.0,
            "coverage_complete": True,
        },
        "code_generation": {
            "target_language": "python",
            "target_framework": "pytest",
            "total_lines_generated": 48,
            "files_written": [
                {"file_path": "tests/test_tickets.py", "relative_path": "tests/test_tickets.py", "lines_count": 48}
            ]
        },
        "tests": [
            {
                "test_key": "TC-SBP101-001",
                "title": "Create ticket with valid payload",
                "scenario_type": "positive",
                "acceptance_criteria_ids": ["AC-01"],
                "target_language": "python",
                "framework": "pytest",
                "status": "PASSED",
                "generated_code": 'def test_create_ticket_valid(client):\n    """AC-01 verify ticket creation"""\n    response = client.post("/api/tickets", json={"title": "Bug", "priority": "high"})\n    assert response.status_code == 201\n    assert response.get_json()["status"] == "OPEN"',
                "request_spec": {"method": "POST", "endpoint": "/api/tickets"},
                "expected_response_spec": {"status_code": 201}
            },
            {
                "test_key": "TC-SBP101-002",
                "title": "Reject ticket creation when title is missing",
                "scenario_type": "negative",
                "acceptance_criteria_ids": ["AC-02"],
                "target_language": "python",
                "framework": "pytest",
                "status": "PASSED",
                "generated_code": 'def test_create_ticket_missing_title(client):\n    """AC-02 verify mandatory title rejection"""\n    response = client.post("/api/tickets", json={"priority": "high"})\n    assert response.status_code == 400\n    assert "title is required" in str(response.get_json())',
                "request_spec": {"method": "POST", "endpoint": "/api/tickets"},
                "expected_response_spec": {"status_code": 400}
            }
        ],
        "unit_tests": {
            "total": 2,
            "passed": 2,
            "failed": 0,
        },
        "results": [
            {
                "method": "POST",
                "endpoint": "/api/tickets",
                "url": "http://localhost:5001/api/tickets",
                "status_code": 201,
                "passed": True,
                "duration_ms": 14,
                "ac_key": "AC-01",
                "test_case_id": "TC-SBP101-001",
                "request_payload": {"title": "Test Ticket", "priority": "high"},
                "response_payload": {"id": "t-101", "title": "Test Ticket", "priority": "high", "status": "OPEN"},
                "assertions": [
                    {"name": "Status code is 201", "passed": True},
                    {"name": "Status is OPEN", "passed": True}
                ]
            }
        ],
        "deviations": [],
    }

    # 1. Test DOCX generation
    docx_path = generate_docx_evidence(evidence_data, out_dir=out_dir)
    assert os.path.isfile(docx_path)

    doc = docx.Document(docx_path)
    full_text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])

    # Assert Coverage is present in DOCX
    assert "Code Coverage" in full_text
    assert "Acceptance Criteria Code Coverage Matrix" in full_text
    assert "AC-01" in full_text
    assert "AC-02" in full_text
    assert "YES [Covered]" in full_text

    # Assert Synthesized Unit Test Raw Code is removed per client requirement
    assert "Synthesized Production Unit Test Code" not in full_text
    assert "def test_create_ticket_valid(client):" not in full_text
    assert "def test_create_ticket_missing_title(client):" not in full_text

    # Assert Test Suite Artifact Reference is present
    assert "Verified Test Suite Artifact" in full_text
    assert "tests/test_tickets.py" in full_text

    # Assert API Testing Screenshots & Highlighted Payloads are present in DOCX
    assert "POSTMAN API EVIDENCE SNAPSHOT" in full_text
    assert "REQUEST PAYLOAD" in full_text
    assert "LIVE CAPTURED RESPONSE" in full_text
    assert "Test Ticket" in full_text

    # 2. Test HTML report generation
    html_path = render_autonomous_evidence_html(evidence_data, out_dir=out_dir)
    assert os.path.isfile(html_path)

    with open(html_path, "r", encoding="utf-8") as f:
        html_content = f.read()

    # Assert Coverage is present in HTML
    assert "Spec. AC Coverage" in html_content
    assert "Acceptance Criteria Coverage Matrix (Specification Coverage)" in html_content
    assert "AC-01" in html_content
    assert "AC-02" in html_content
    assert "YES [Covered]" in html_content

    # Assert Synthesized Unit Test Raw Code is removed from HTML per client requirement
    assert "Synthesized Production Unit Test Code Artifacts" not in html_content
    assert "def test_create_ticket_valid(client):" not in html_content

    # Assert Test Suite Artifact Reference is present in HTML
    assert "Verified Test Suite Artifact" in html_content

    # Assert API Testing Snapshots & Payloads are present in HTML
    assert "POSTMAN API EVIDENCE SNAPSHOT" in html_content
    assert "REQUEST PAYLOAD (BODY SENT)" in html_content
    assert "LIVE CAPTURED SERVER RESPONSE" in html_content
