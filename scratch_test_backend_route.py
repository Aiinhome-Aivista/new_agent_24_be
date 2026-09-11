import requests
import json
import os

BASE_API = "http://localhost:5000/api/v1"
TARGET_API = "http://localhost:5001"

s = requests.Session()
r_login = s.post(f"{BASE_API}/login", json={"email": "developer@tdd.local", "password": "Passw0rd!"})
login_data = r_login.json()
token = login_data.get("data", {}).get("access_token")
print(f"Login status: {r_login.status_code}, Token obtained: {bool(token)}")

headers = {"Authorization": f"Bearer {token}"} if token else {}
r_run = s.post(
    f"{BASE_API}/api-executor/autonomous-run",
    headers=headers,
    json={"base_url": TARGET_API}
)
print(f"Autonomous run response code: {r_run.status_code}")
resp = r_run.json()
data = resp.get("data", {})

print(f"Evidence Key: {data.get('evidence_key')}")
print(f"Total Endpoints Executed: {data.get('total_endpoints')}")
print(f"Passed Endpoints: {data.get('passed_endpoints')}")
print(f"Total Deviations: {data.get('total_deviations')}")
print(f"Summary Recommendation: {data.get('summary_recommendation')}")
print(f"Docx path: {data.get('docx_path')}")
print(f"Workspace Docx path: {data.get('workspace_docx_path')}")

print("\n--- Executed Scenarios ---")
for idx, res in enumerate(data.get("results", [])):
    print(f"{idx+1}. [{res.get('method')}] {res.get('test_key')} -> Status {res.get('status_code')} (Passed: {res.get('passed')})")

print("\n--- Workspace Docx Verification ---")
ws_docx = data.get("workspace_docx_path")
if ws_docx and os.path.exists(ws_docx):
    print(f"Found docx in workspace: {ws_docx} (Size: {os.path.getsize(ws_docx):,} bytes)")
    import docx
    d = docx.Document(ws_docx)
    tables_count = len(d.tables)
    images_count = sum(1 for p in d.paragraphs if any('graphic' in run._r.xml for run in p.runs))
    print(f"Document paragraphs: {len(d.paragraphs)}, tables: {tables_count}, images: {images_count}")
